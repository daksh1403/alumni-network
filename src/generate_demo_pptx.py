#!/usr/bin/env python3
"""Generate the DA2 Demo Presentation (PPTX) for the Alumni Network DBMS project."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "report", "Alumni_Network_Demo_Presentation.pptx")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x40, 0x40, 0x40)
LIGHT = RGBColor(0xEE, 0xF3, 0xF9)

prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def add_slide(title=None):
    s = prs.slides.add_slide(BLANK)
    if title:
        box = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
        tf = box.text_frame
        tf.text = title
        r = tf.paragraphs[0].runs[0]
        r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = ACCENT
    return s

def bullets(slide, items, top=1.5, size=20, left=0.8, width=11.7):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(7.2 - top))
    tf = box.text_frame; tf.word_wrap = True
    first = True
    for it in items:
        lvl, text = ((it[1], it[0]) if isinstance(it, tuple) else (0, it))
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = text; p.level = lvl
        f = p.runs[0].font
        f.size = Pt(size - lvl * 2); f.color.rgb = GREY
    return box

def codebox(slide, text, top=4.4, left=0.8, width=11.7, height=2.6, size=12):
    shape = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = shape.text_frame; tf.word_wrap = False
    shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT
    lines = text.strip("\n").split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln if ln else " "
        p.runs[0].font.name = "Courier New"; p.runs[0].font.size = Pt(size)
    return shape

# ---------------------------------------------------------------- 1 title
s = prs.slides.add_slide(BLANK)
box = s.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(11), Inches(2.4))
tf = box.text_frame
tf.text = "ALUMNI NETWORK AND ENGAGEMENT PLATFORM"
tf.paragraphs[0].runs[0].font.size = Pt(44)
tf.paragraphs[0].runs[0].font.bold = True
tf.paragraphs[0].runs[0].font.color.rgb = ACCENT
p = tf.add_paragraph(); p.text = "DA2 - Database Implementation using SQL & PL/SQL"
p.runs[0].font.size = Pt(24); p.runs[0].font.color.rgb = GREY
p = tf.add_paragraph()
p.text = "\nSagarika Kaistha - 25BCE5091   |   Praveen G - 25BCE5092   |   Daksh Agarwal - 25BCE5098"
p.runs[0].font.size = Pt(16); p.runs[0].font.color.rgb = GREY

# ---------------------------------------------------------------- 2 agenda
s = add_slide("Agenda")
bullets(s, [
    "From design to working database - implementation pipeline",
    "Schema implementation: tables, constraints, indexes",
    "PL/SQL layer: sequences, triggers, procedures, functions, package",
    "Verification: every rule tested against the live database",
    "Live demonstration walkthrough",
    "Conclusion and Q&A",
], size=24)

# ---------------------------------------------------------------- 3 pipeline
s = add_slide("Implementation Pipeline")
steps = ["01_schema.sql\nDDL + constraints", "02_plsql.sql\nPL/SQL objects",
         "03_seed_data.sql\nsample data", "04_demo_queries.sql\nverification suite"]
x = 0.75
for i, txt in enumerate(steps):
    shape = s.shapes.add_textbox(Inches(x), Inches(2.6), Inches(2.7), Inches(1.6))
    tf = shape.text_frame; tf.word_wrap = True
    tf.text = txt.split("\n")[0]
    r = tf.paragraphs[0].runs[0]; r.font.bold = True; r.font.size = Pt(18); r.font.color.rgb = ACCENT
    p = tf.add_paragraph(); p.text = txt.split("\n")[1]
    p.runs[0].font.size = Pt(14); p.runs[0].font.color.rgb = GREY
    shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT
    x += 3.05
codebox(s, """sqlplus alumni/alumni123@localhost:1521/FREEPDB1
@sql/01_schema.sql && @sql/02_plsql.sql && @sql/03_seed_data.sql && @sql/04_demo_queries.sql
# zero ORA errors on clean build - logs/run_*.log""", top=5.0)

# ---------------------------------------------------------------- 4 schema stats
s = add_slide("What Was Built")
bullets(s, [
    "22 tables - supertype PERSON with ALUMNI / STUDENT / ADMIN subtypes",
    "30 foreign keys, 22 primary keys (composite on junction tables), 12 UNIQUE, 65 CHECK constraints",
    "61 indexes incl. all 18 recommended in the DA1 design document",
    "3 design views: AlumniProfileView, EventSummaryView, DonationSummaryView",
    "Sample data: 15 persons, 10 alumni, 5 events, 12 donations, 6 jobs, 4 mentorships, full forum thread",
], size=22)

# ---------------------------------------------------------------- 5 type mapping
s = add_slide("Design -> Oracle Mapping Decisions")
bullets(s, [
    ("BOOLEAN -> NUMBER(1) CHECK IN (0,1)", 0),
    ("TEXT -> CLOB; VARCHAR -> VARCHAR2", 0),
    ("AUTO_INCREMENT -> SEQUENCE + BEFORE INSERT trigger (17 key triggers)", 0),
    ("DATETIME DEFAULT NOW() -> TIMESTAMP DEFAULT SYSTIMESTAMP", 0),
    ("TIME -> VARCHAR2(5) with REGEXP 'HH24:MI' check (no TIME type in Oracle)", 0),
    ("ON UPDATE CASCADE omitted - not supported by Oracle (documented)", 0),
    ("COMMENT table renamed COMMENT_TABLE (SQL*Plus keyword clash)", 0),
], size=22)

# ---------------------------------------------------------------- 6 sequences/triggers
s = add_slide("Auto-Increment via Sequences + Triggers")
codebox(s, """CREATE SEQUENCE seq_person START WITH 1 INCREMENT BY 1 NOCACHE;

CREATE OR REPLACE TRIGGER trg_person_bi BEFORE INSERT ON PERSON
FOR EACH ROW
BEGIN
    :NEW.PersonID := NVL(:NEW.PersonID, seq_person.NEXTVAL);
END;
/
-- INSERT INTO PERSON (...) VALUES (...);  -- ID assigned automatically""",
       top=1.6, height=3.2, size=14)
bullets(s, ["18 NOCACHE sequences -> deterministic IDs for seed data & demos",
            "INSERT statements never need to supply surrogate keys"],
        top=5.2, size=18)

# ---------------------------------------------------------------- 7 business rules
s = add_slide("Business-Rule Triggers")
bullets(s, [
    ("R1  TRG_DONATION_BUSINESS     no future donations; auto receipt RCP-<seq>-<yyyy>      ORA-20001", 0),
    ("R2  TRG_EVENTREG_BUSINESS     capacity enforcement + no late registrations             ORA-20002/-20005", 0),
    ("R3  TRG_MENTORSHIP_BUSINESS   mentor <> mentee; end >= start; auto-complete           ORA-20003/-20006", 0),
    ("R4  TRG_JOB_BUSINESS          expiry must be after posting date                        ORA-20004", 0),
    ("R5  TRG_ALUMNI_AUDIT          AFTER row trigger writes full audit trail to ALUMNI_AUDIT", 0),
], size=19, top=1.6)
codebox(s, """IF v_booked >= v_capacity THEN
    RAISE_APPLICATION_ERROR(-20002,
        'Event is FULL (' || v_booked || '/' || v_capacity ||
        ' seats taken). Registration rejected.');
END IF;""", top=4.9, height=1.9, size=13)

# ---------------------------------------------------------------- 8 procedures
s = add_slide("Stored Procedures")
bullets(s, [
    ("sp_register_event(alumni, event)      duplicate-safe event registration", 0),
    ("sp_record_donation(donor, amt, ...)   OUT receipt generated via RETURNING INTO", 0),
    ("sp_apply_to_job(job, applicant)       rejects closed / expired / duplicate applications", 0),
    ("sp_close_mentorship(id, end, fb, rating)  FOR UPDATE lock; ACTIVE-only closable", 0),
], size=21, top=1.6)
codebox(s, """SQL> EXEC sp_record_donation(6, 5000, 'Online', 'General Fund', :v);
[OK] Donation of Rs.5000.00 recorded from Ananya Iyer. Receipt: RCP-1012-2026""",
       top=4.6, height=1.7, size=14)

# ---------------------------------------------------------------- 9 functions+package
s = add_slide("Functions & Package PKG_ALUMNI_REPORTS")
bullets(s, [
    ("fn_total_donations(person_id)   lifetime donation total", 0),
    ("fn_event_occupancy(event_id)    occupancy % (NULL if unlimited)", 0),
    ("fn_active_mentees(mentor_id)    active mentee count", 0),
    ("pkg.top_donors(n)               explicit cursor + RANK() analytical", 0),
    ("pkg.dept_alumni_summary         department-wise statistics", 0),
    ("pkg.get_alumni_skills(id)       skill list via cursor loop", 0),
], size=20, top=1.6)
codebox(s, """---- TOP 5 DONORS ----
RANK  ID    NAME                AMOUNT
1     1     Rahul Sharma        75000.00
1     2     Priya Menon         75000.00
3     10    Meera Krishnan      60000.00""", top=5.0, height=2.0, size=13)

# ---------------------------------------------------------------- 10 verification
s = add_slide("Verification - Every Rule Tested Live")
codebox(s, """CAUGHT: ORA-20001: Donation date cannot be in the future: 31-AUG-2026
CAUGHT: ORA-20002: Event is FULL (3/3 seats taken). Registration rejected.
CAUGHT: ORA-20003: Invalid mentorship: MentorID and MenteeID are the same person.
CAUGHT: ORA-20004: Job expiry date must be after the posting date.
CAUGHT: ORA-20205: You have already applied to "Senior Backend Developer".

[OK] Alumnus #6 registered for "Annual Tech Meetup" (#5)
[OK] Application submitted for "Product Management Intern" (#2)
[OK] Mentorship #2 marked Completed (rating: 5/5).""",
       top=1.6, height=3.6, size=14)
bullets(s, ["Clean build: zero errors; all 100+ objects VALID in USER_OBJECTS",
            "Evidence captured in logs/run_01...04 spool files"],
        top=5.5, size=19)

# ---------------------------------------------------------------- 11 SQL showcase
s = add_slide("Demonstration Queries - SQL Feature Coverage")
bullets(s, [
    ("Joins (inner/left, 3-4 tables) + LISTAGG aggregation                    Q1-Q2", 0),
    ("GROUP BY + HAVING; average-comparison subquery                          Q3-Q4", 0),
    ("Correlated NOT EXISTS                                                   Q5", 0),
    ("Analytical RANK() over registrations                                    Q6", 0),
    ("CONNECT BY PRIOR over self-referencing comment tree                     Q7", 0),
    ("Views | SAVEPOINT/ROLLBACK | functions/package calls | audit trail      Q8-Q14", 0),
], size=22, top=1.7)

# ---------------------------------------------------------------- 12 live demo flow
s = add_slide("Live Demo Plan")
bullets(s, [
    "1. Show schema inventory (USER_TABLES / row counts)",
    "2. Run Q1-Q7 queries - joins to hierarchical comments",
    "3. Register alumnus for an event via procedure; then attempt one more seat on FULL reunion",
    "4. Record donation -> show generated receipt number",
    "5. Duplicate job application -> friendly rejection",
    "6. Close a mentorship with rating; query MENTORSHIP",
    "7. Show ALUMNI_AUDIT rows written automatically",
], size=22, top=1.6)

# ---------------------------------------------------------------- 13 conclusion
s = add_slide("Conclusion")
bullets(s, [
    "Complete DA1 design implemented as a running Oracle database",
    "Integrity enforced at two levels: declarative constraints + PL/SQL business rules",
    "Fully scripted and reproducible (works on Live SQL / lab Oracle 11g+)",
    "Future work: application front-end, partitioning, roles, REST API over package",
], size=23)

# ---------------------------------------------------------------- 14 thanks
s = prs.slides.add_slide(BLANK)
box = s.shapes.add_textbox(Inches(1.2), Inches(3.0), Inches(11), Inches(1.5))
tf = box.text_frame; tf.text = "Thank You - Questions?"
r = tf.paragraphs[0].runs[0]
r.font.size = Pt(48); r.font.bold = True; r.font.color.rgb = ACCENT

os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print("Saved:", OUT)

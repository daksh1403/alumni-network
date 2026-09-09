#!/usr/bin/env python3
"""Generate the DA2 Implementation Report (DOCX) for the Alumni Network DBMS project."""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "report", "Alumni_Network_Implementation.docx")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)

def h1(t):
    p = doc.add_heading(t, level=1)
    for r in p.runs: r.font.color.rgb = ACCENT
    return p

def h2(t):
    p = doc.add_heading(t, level=2)
    for r in p.runs: r.font.color.rgb = ACCENT
    return p

def para(t, bold=False, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t); r.bold = bold; r.italic = italic
    return p

def bullet(t):
    return doc.add_paragraph(t, style="List Bullet")

def code(lines):
    for ln in lines.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.25)
        r = p.add_run(ln if ln else " ")
        r.font.name = "Consolas"; r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x20, 0x20, 0x20)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, htxt in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt); r.bold = True; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val)); r.font.size = Pt(9.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t

# ------------------------------------------------------------------ title
for _ in range(4): doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("ALUMNI NETWORK AND\nENGAGEMENT PLATFORM"); r.bold = True; r.font.size = Pt(28); r.font.color.rgb = ACCENT
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Database Management Systems Project"); r.font.size = Pt(14)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("DA2 - Database Implementation using SQL and PL/SQL\nwith Report Submission & Project Demonstration")
r.font.size = Pt(13); r.bold = True
doc.add_paragraph(); 
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Team Members:\nSagarika Kaistha - 25BCE5091\nPraveen G - 25BCE5092\nDaksh Agarwal - 25BCE5098")
r.font.size = Pt(12)
doc.add_page_break()

# ------------------------------------------------------------------ toc
h1("TABLE OF CONTENTS")
for item in [
    "1. Introduction", "2. Implementation Environment",
    "3. Schema Implementation (SQL DDL)", "4. PL/SQL Implementation",
    "5. Views", "6. Verification and Testing",
    "7. Demonstration Queries", "8. Conclusion",
    "Appendix A - How to Run", "Appendix B - Script Inventory",
]:
    doc.add_paragraph(item)

# ------------------------------------------------------------------ 1
h1("1. Introduction")
para("In the previous deliverable (DA1), the Alumni Network and Engagement Platform was "
     "designed using the ER/EER model, normalized up to BCNF/4NF, and mapped to a relational "
     "schema of 21 relations. This report documents the implementation of that design as a "
     "working Oracle database.")
para("The implementation comprises:")
bullet("22 tables implementing every entity, supertype/subtype hierarchy, lookup table and "
       "junction table from the design, with all primary-key, foreign-key, UNIQUE and CHECK "
       "constraints enforced by the DBMS.")
bullet("18 sequences and 17 auto-increment triggers providing Oracle-style surrogate key generation.")
bullet("5 business-rule triggers enforcing integrity that table constraints cannot express "
       "(capacity limits, date logic, audit trail).")
bullet("4 stored procedures, 3 functions and 1 package (PKG_ALUMNI_REPORTS) encapsulating the "
       "core platform operations: event registration, donation recording with receipt "
       "generation, job applications and mentorship lifecycle.")
bullet("3 views defined in the DA1 design document (AlumniProfileView, EventSummaryView, "
       "DonationSummaryView).")
bullet("A demonstration suite of 14 query scenarios covering joins, aggregates, subqueries, "
       "analytical functions, hierarchical queries and transaction control.")

# ------------------------------------------------------------------ 2
h1("2. Implementation Environment")
table(["Component", "Details"],
      [["DBMS", "Oracle AI Database 26ai Free (23.26.2.0.0), pluggable database FREEPDB1"],
       ["SQL client", "SQL*Plus Release 23.26.1.0.0 (Oracle Instant Client)"],
       ["Hosting", "Docker container gvenzl/oracle-free on macOS; port 1521"],
       ["Project schema", "User ALUMNI with CREATE TABLE / VIEW / PROCEDURE / TRIGGER / SEQUENCE privileges"],
       ["Scripts", "sql/00_drop_all.sql ... sql/run_all.sql executed in order via SQL*Plus @"]],
      widths=[1.6, 4.9])
para("Every script is plain Oracle SQL/PL-SQL with no proprietary client features, so the same "
     "files run unchanged on Oracle Live SQL or any college lab Oracle 11g+ installation.",
     italic=True)

# ------------------------------------------------------------------ 3
h1("3. Schema Implementation (SQL DDL)")
h2("3.1 Data Type Mapping")
para("The DA1 schema was written against generic SQL types. The mapping used for Oracle:")
table(["Design type", "Oracle type", "Notes"],
      [["INT / BIGINT", "NUMBER(p)", "precision per attribute"],
       ["DECIMAL(m,n)", "NUMBER(m,n)", "e.g. DONATION.Amount NUMBER(12,2)"],
       ["VARCHAR(n)", "VARCHAR2(n)", "Oracle-recommended variable text"],
       ["TEXT", "CLOB", "descriptions, posts, comments, feedback"],
       ["BOOLEAN", "NUMBER(1) + CHECK IN (0,1)", "IsActive, IsPinned, IsAnonymous"],
       ["DATETIME DEFAULT NOW()", "TIMESTAMP DEFAULT SYSTIMESTAMP", "CreatedAt, PostedDate, ..."],
       ["AUTO_INCREMENT", "SEQUENCE + BEFORE INSERT trigger", "see section 4.1"],
       ["TIME", "VARCHAR2(5) + REGEXP CHECK 'HH24:MI'", "EVENT.EventTime (Oracle has no TIME type)"]],
      widths=[1.7, 2.4, 2.4])

h2("3.2 Tables Implemented and Sample Data Volume")
rows = [
    ["PERSON", "Supertype (EER generalization)", 15],
    ["ALUMNI / STUDENT / ADMIN", "Subtypes, PK = FK to PERSON", "10 / 3 / 2"],
    ["DEPARTMENT / BATCH", "Academic structure", "4 / 5"],
    ["COMPANY / COMPANY_SIZE", "Employers + size lookup", "8 / 5"],
    ["SKILL / ALUMNI_SKILL", "Skill catalogue + M:N junction", "8 / 14"],
    ["EVENT_TYPE / EVENT", "Event lookup + events", "4 / 5"],
    ["EVENT_REGISTRATION", "M:N alumni-event with attendance status", 11],
    ["DONATION", "Contributions; auto receipt numbers", 12],
    ["JOB / JOB_APPLICATION", "Job portal + M:N applications", "6 / 8"],
    ["MENTORSHIP_AREA / MENTORSHIP", "Mentorship domains + engagements", "5 / 4"],
    ["FORUM / POST / COMMENT_TABLE", "Discussion hierarchy, self-referencing replies", "3 / 6 / 8"],
    ["ALUMNI_AUDIT", "Trigger-populated audit trail", 10],
]
table(["Table(s)", "Role in design", "Rows"], rows, widths=[2.2, 3.3, 0.9])
para("COMMENT_TABLE implements the designed COMMENT relation; the name is suffixed because some "
     "SQL*Plus configurations treat bare COMMENT ambiguously.", italic=True)

h2("3.3 Constraint Enforcement (verified via USER_CONSTRAINTS)")
table(["Type", "Count", "Examples"],
      [["PRIMARY KEY", "22", "one per table; composite PKs on the three junction tables"],
       ["FOREIGN KEY", "30", "subtype FKs ON DELETE CASCADE, HOD/company SET NULL, self-referencing comment parent"],
       ["UNIQUE", "12", "Email, StudentID, ReceiptNumber, TransactionID, (Mentor,Mentee,Start), (Job,Applicant)"],
       ["CHECK", "65", "gender, CGPA range, rating 1-5, payment methods, attendance status, HH24:MI time format"]],
      widths=[1.3, 0.8, 4.3])
para("All 18 indexes recommended by the design document were created for frequent join/filter paths.")

# ------------------------------------------------------------------ 4
h1("4. PL/SQL Implementation")
h2("4.1 Sequences and Auto-Increment Triggers")
para("Each surrogate key has a NOCACHE sequence consumed by a BEFORE INSERT row trigger, so INSERT "
     "statements may omit IDs entirely (the pattern used throughout seed data):")
code("""CREATE SEQUENCE seq_person START WITH 1 INCREMENT BY 1 NOCACHE;

CREATE OR REPLACE TRIGGER trg_person_bi BEFORE INSERT ON PERSON
FOR EACH ROW
BEGIN
    :NEW.PersonID := NVL(:NEW.PersonID, seq_person.NEXTVAL);
END;
/""")
para("Inventory: 18 sequences (including seq_receipt for receipt numbers) and 17 auto-increment triggers.")

h2("4.2 Business-Rule Triggers")
para("Five triggers enforce rules that column constraints cannot express. Each rule raises a "
     "distinct application error (RAISE_APPLICATION_ERROR) so failures are self-explanatory:")
table(["#", "Rule enforced", "Trigger", "Error", "Verified result"],
      [["R1", "Donation cannot be future-dated; missing receipts auto-generated as RCP-<seq>-<yyyy>",
        "TRG_DONATION_BUSINESS", "ORA-20001", "rejected, receipt generated"],
       ["R2", "Event registration rejected when capacity is full or event already conducted",
        "TRG_EVENTREG_BUSINESS", "ORA-20002 / ORA-20005", "'FULL (3/3)' observed"],
       ["R3", "Mentor <> mentee; end date >= start date; setting end date completes an active engagement",
        "TRG_MENTORSHIP_BUSINESS", "ORA-20003 / ORA-20006", "self-mentoring blocked"],
       ["R4", "JOB.ExpiryDate must be later than posting date", "TRG_JOB_BUSINESS", "ORA-20004", "rejected"],
       ["R5", "Every ALUMNI insert/update/delete logged to ALUMNI_AUDIT with user and timestamp",
        "TRG_ALUMNI_AUDIT (AFTER)", "-", "10 rows captured"]],
      widths=[0.4, 2.6, 1.5, 1.1, 1.4])
code("""CREATE OR REPLACE TRIGGER trg_eventreg_business
BEFORE INSERT OR UPDATE ON EVENT_REGISTRATION
FOR EACH ROW
DECLARE
    v_capacity EVENT.MaxCapacity%TYPE;
    v_booked   NUMBER;
    v_edate    EVENT.EventDate%TYPE;
BEGIN
    SELECT NVL(MaxCapacity, 999999), EventDate
      INTO v_capacity, v_edate
      FROM EVENT WHERE EventID = :NEW.EventID;

    IF TRUNC(SYSDATE) > v_edate THEN
        RAISE_APPLICATION_ERROR(-20005,
            'Cannot register: event already conducted.');
    END IF;

    SELECT COUNT(*) INTO v_booked
      FROM EVENT_REGISTRATION
     WHERE EventID = :NEW.EventID
       AND AttendanceStatus IN ('Registered', 'Attended');

    IF v_booked >= v_capacity THEN
        RAISE_APPLICATION_ERROR(-20002,
            'Event is FULL (' || v_booked || '/' || v_capacity ||
            ' seats taken). Registration rejected.');
    END IF;
END;
/""")

h2("4.3 Stored Procedures")
table(["Procedure", "Purpose", "Key behaviours"],
      [["sp_register_event(alumni, event)", "Register for an event",
        "duplicate-safe (DUP_VAL_ON_INDEX handled), capacity rule inherited from trigger"],
       ["sp_record_donation(donor, amount, method, purpose, OUT receipt)",
        "Record a contribution", "validates donor, generates receipt via RETURNING INTO, commits"],
       ["sp_apply_to_job(job, applicant, resume)", "Apply to a posting",
        "rejects closed/expired jobs and duplicate applications (ORA-20203..20206)"],
       ["sp_close_mentorship(id, end_date, feedback, rating)", "Complete a mentorship",
        "row-locked with FOR UPDATE, only ACTIVE engagements closable (ORA-20207)"]],
      widths=[2.1, 1.6, 2.8])

h2("4.4 Functions")
table(["Function", "Returns"],
      [["fn_total_donations(person_id)", "Lifetime donation total of an alumnus"],
       ["fn_event_occupancy(event_id)", "Occupancy percentage (NULL when capacity unlimited)"],
       ["fn_active_mentees(mentor_id)", "Number of active mentees guided by an alumnus"]],
      widths=[2.6, 3.9])

h2("4.5 Package PKG_ALUMNI_REPORTS")
para("The package separates specification from body and demonstrates explicit cursors, cursor FOR "
     "loops and analytical SQL inside PL/SQL:")
bullet("top_donors(n) - RANK() over aggregated donations, printed via DBMS_OUTPUT")
bullet("dept_alumni_summary - department-wise counts and average graduation year")
bullet("get_alumni_skills(person_id) - returns a comma-separated skill list built with a cursor loop")

# ------------------------------------------------------------------ 5
h1("5. Views")
para("The three views specified in the DA1 design were created verbatim (BOOLEAN adapted to NUMBER(1)):")
code("""CREATE OR REPLACE VIEW AlumniProfileView AS
SELECT a.PersonID, p.FirstName, p.LastName, p.Email, a.GraduationYear,
       d.DeptName, c.CompanyName, a.CurrentPosition
  FROM ALUMNI a
  JOIN PERSON p     ON a.PersonID = p.PersonID
  JOIN DEPARTMENT d ON a.DeptID   = d.DeptID
  LEFT JOIN COMPANY c ON a.CurrentCompanyID = c.CompanyID
 WHERE a.IsActive = 1;""")
bullet("EventSummaryView - registrations per event with organiser name and capacity")
bullet("DonationSummaryView - count, sum and last donation date per alumnus")

# ------------------------------------------------------------------ 6
h1("6. Verification and Testing")
h2("6.1 Build verification")
para("Executing 00 -> 01 -> 02 -> 03 on a clean schema produced zero ORA errors: 22 tables, 61 indexes, "
     "18 sequences, 22 triggers, 3 views, 4 procedures, 3 functions and 1 package compiled VALID "
     "(confirmed through USER_OBJECTS / USER_ERRORS).")
h2("6.2 Business-rule test results (from spool log run_04_demo_queries.log)")
code("""CAUGHT: ORA-20001: Donation date cannot be in the future: 31-AUG-2026
CAUGHT: ORA-20002: Event is FULL (3/3 seats taken). Registration rejected.
CAUGHT: ORA-20003: Invalid mentorship: MentorID and MenteeID are the same person.
CAUGHT: ORA-20004: Job expiry date must be after the posting date.
CAUGHT: ORA-20205: You have already applied to "Senior Backend Developer".

[OK] Donation of Rs.5000.00 recorded from Ananya Iyer. Receipt: RCP-1012-2026
[OK] Alumnus #6 registered for "Annual Tech Meetup" (#5)
[OK] Application submitted for "Product Management Intern" (#2)
[OK] Mentorship #2 marked Completed (rating: 5/5).""")
h2("6.3 Package output sample")
code("""---- TOP 5 DONORS ----
RANK  ID    NAME                        AMOUNT
1     1     Rahul Sharma                75000.00
1     2     Priya Menon                 75000.00
3     10    Meera Krishnan              60000.00
4     3     Arjun Patel                 55000.00
5     9     Rohit Verma                 35000.00""")
h2("6.4 Audit trail evidence")
code("""AUDIT_ID  OPERATION  NEW_VALUES
1         INSERT     Year=2019, Dept=1, Pos=Senior Software Engineer
...
10        INSERT     Year=2022, Dept=4, Pos=Founder - AgriTech Startup""")
para("(Screenshot placeholders: insert live screenshots of the spool logs during the viva demo.)",
     italic=True)

# ------------------------------------------------------------------ 7
h1("7. Demonstration Queries")
table(["Q#", "SQL concept showcased"],
      [["Q1-Q2", "Inner / LEFT OUTER joins across 3-4 tables; LISTAGG aggregation"],
       ["Q3-Q4", "GROUP BY + HAVING; subquery comparing against average"],
       ["Q5", "Correlated NOT EXISTS (alumni who never donated)"],
       ["Q6", "Analytical RANK() over event popularity"],
       ["Q7", "Hierarchical CONNECT BY PRIOR over the self-referencing comment tree"],
       ["Q8", "The three design views"],
       ["Q9", "Transaction control: SAVEPOINT / ROLLBACK TO SAVEPOINT with a CHECK violation"],
       ["Q10-Q11", "Functions and package invocation from SQL and SQL*Plus EXEC"],
       ["Q12-Q13", "Procedures: happy paths plus each business-rule rejection"],
       ["Q14", "Audit trail populated by trigger R5"]],
      widths=[0.9, 5.6])

# ------------------------------------------------------------------ 8
h1("8. Conclusion")
para("The complete DA1 design is now a running Oracle database. Every entity, relationship and "
     "constraint from the ER/EER model is enforced physically; PL/SQL adds the procedural layer "
     "(key generation, business validation, auditing, reporting API). The build is fully scripted "
     "and reproducible, and the demonstration suite shows both successful flows and deliberate "
     "violations being caught by the database itself.")
para("Future work: add a web/application front-end, partition DONATION by year, introduce fine-grained "
     "roles, and expose the package as REST endpoints for a mobile app.", italic=True)

# ------------------------------------------------------------------ appendix A
doc.add_page_break()
h1("Appendix A - How to Run")
code("""# Local SQL*Plus
cd <project-root>
sqlplus system/<password>@localhost:1521/FREEPDB1
  CREATE USER alumni IDENTIFIED BY alumni123 QUOTA UNLIMITED ON USERS;
  GRANT CREATE SESSION, CREATE TABLE, CREATE VIEW, CREATE PROCEDURE,
        CREATE TRIGGER, CREATE SEQUENCE TO alumni;

CONNECT alumni/alumni123@localhost:1521/FREEPDB1
@sql/01_schema.sql
@sql/02_plsql.sql
@sql/03_seed_data.sql
SET SERVEROUTPUT ON
@sql/04_demo_queries.sql""")
para("On Oracle Live SQL: paste the contents of 01 -> 02 -> 03 -> 04 in order into the worksheet "
     "and enable server output in settings.")

h1("Appendix B - Script Inventory")
table(["File", "Contents"],
      [["sql/00_drop_all.sql", "Tears down views, PL/SQL objects, triggers, tables, sequences"],
       ["sql/01_schema.sql", "22 tables, constraints, 18 indexes, audit table"],
       ["sql/02_plsql.sql", "Sequences, auto-increment + business triggers, views, procedures, functions, package"],
       ["sql/03_seed_data.sql", "Sample data + row-count sanity check"],
       ["sql/04_demo_queries.sql", "14 demonstration scenarios"],
       ["sql/run_all.sql", "Master runner producing logs/run_*.log spools"],
       ["logs/", "Captured execution logs used as report evidence"]],
      widths=[2.0, 4.5])

os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print("Saved:", OUT)

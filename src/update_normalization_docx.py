#!/usr/bin/env python3
"""
Update Alumni_Network_Normalization_Corrected.docx in place:
  1. Add a "Sample Data" table (5 records) to every relation section.
  2. Exclude the COMPANY-JOB ("at") relationship from the JOB relation.
  3. Fix the relation count 13 -> 14 (add STUDENT_EMAIL to the summary).
"""

import os
import shutil
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOCX = "/Users/dakshagarwal/dbms-project/Alumni_Network_Normalization_Corrected.docx"
BACKUP = "/Users/dakshagarwal/dbms-project/Alumni_Network_Normalization_Corrected.bak.docx"

HEADER_FILL = "1F4E78"
ROW_FILL = "EEF5FB"


def set_cell_shading(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def style_cell(cell, text, bold=False, color=None, font_size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(str(text))
    run.font.name = "Calibri"
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def make_sample_table(headers, rows, font_size=10, col_widths=None):
    """Return a styled Table Grid table (appended to the doc body)."""
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if col_widths:
        # Set fixed layout so wide tables fit the page
        tblPr = table._tbl.tblPr
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tblPr.append(layout)
        tblW = tblPr.find(qn("w:tblW"))
        if tblW is None:
            tblW = OxmlElement("w:tblW")
            tblPr.append(tblW)
        total = sum(col_widths)
        tblW.set(qn("w:w"), str(total))
        tblW.set(qn("w:type"), "dxa")
        # Set each grid column width
        grid = table._tbl.find(qn("w:tblGrid"))
        for gc, w in zip(grid.findall(qn("w:gridCol")), col_widths):
            gc.set(qn("w:w"), str(w))
        # And each cell width
        for row in table.rows:
            for cell, w in zip(row.cells, col_widths):
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.append(tcW)
                tcW.set(qn("w:w"), str(w))
                tcW.set(qn("w:type"), "dxa")

    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        set_cell_shading(cell, HEADER_FILL)
        style_cell(cell, h, bold=True, color="FFFFFF", font_size=font_size)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            if i % 2 == 0:
                set_cell_shading(cell, ROW_FILL)
            style_cell(cell, val, font_size=font_size)

    return table


def insert_after(anchor_p, caption_text, headers, rows, font_size=10, col_widths=None):
    """Insert a caption + sample-data table right after the anchor paragraph."""
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = cap.add_run(caption_text)
    run.font.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(10)

    tbl = make_sample_table(headers, rows, font_size, col_widths)

    anchor_p._p.addnext(tbl._tbl)
    anchor_p._p.addnext(cap._p)


def find_paragraph(substring, start=0):
    for i, p in enumerate(doc.paragraphs):
        if i >= start and substring in p.text:
            return i, p
    raise ValueError(f"Paragraph containing {substring!r} not found")


def delete_paragraph(p):
    p._p.getparent().remove(p._p)


# ---------------------------------------------------------------------------
# Sample data definitions (consistent FK values across tables)
# ---------------------------------------------------------------------------
SAMPLE_DATA = [
    # (anchor text, caption, headers, rows, font_size)
    (
        "DeptID \u2192 DeptName, DeptCode, HODName, EstablishedYear",
        "Sample Data \u2014 DEPARTMENT (5 records)",
        ["DeptID", "DeptName", "DeptCode", "HODName", "EstablishedYear"],
        [
            ["1", "Computer Science", "CSE", "Dr. R. Sharma", "1998"],
            ["2", "Electronics & Comm.", "ECE", "Dr. S. Iyer", "1999"],
            ["3", "Mechanical Engg.", "MECH", "Dr. A. Verma", "1995"],
            ["4", "Electrical Engg.", "EEE", "Dr. P. Nair", "2001"],
            ["5", "Civil Engg.", "CIVIL", "Dr. K. Rao", "1994"],
        ],
        10,
    ),
    (
        "BatchID \u2192 BatchYear, Section, TotalStudents, DeptID",
        "Sample Data \u2014 BATCH (5 records)",
        ["BatchID", "BatchYear", "Section", "TotalStudents", "DeptID"],
        [
            ["1", "2023", "A", "120", "1"],
            ["2", "2023", "B", "118", "1"],
            ["3", "2024", "A", "125", "2"],
            ["4", "2024", "B", "122", "3"],
            ["5", "2025", "A", "130", "4"],
        ],
        10,
    ),
    (
        "CompanyID \u2192 CompanyName, Industry, CompanySize, Website, Headquarters",
        "Sample Data \u2014 COMPANY (5 records)",
        ["CompanyID", "CompanyName", "Industry", "CompanySize", "Website", "Headquarters"],
        [
            ["101", "TCS", "IT Services", "500000+", "tcs.com", "Mumbai"],
            ["102", "Infosys", "IT Services", "300000+", "infosys.com", "Bengaluru"],
            ["103", "Google", "Technology", "100000+", "google.com", "Mountain View"],
            ["104", "Reliance", "Conglomerate", "200000+", "reliance.com", "Mumbai"],
            ["105", "Zoho", "Software", "15000", "zoho.com", "Chennai"],
        ],
        10,
    ),
    (
        "SkillID \u2192 SkillName, SkillCategory, Description",
        "Sample Data \u2014 SKILL (5 records)",
        ["SkillID", "SkillName", "SkillCategory", "Description"],
        [
            ["1", "Java", "Programming", "OOP and backend development"],
            ["2", "Python", "Programming", "Data analysis and automation"],
            ["3", "SQL", "Database", "Querying and database design"],
            ["4", "Communication", "Soft Skill", "Verbal and written communication"],
            ["5", "Project Management", "Management", "Planning and team coordination"],
        ],
        10,
    ),
    (
        "AlumniID \u2192 FirstName, LastName, Email, DateOfBirth, Gender, Address_City, Address_State, Address_PinCode, GraduationYear, DeptID, BatchID, CompanyID, CurrentPosition, LinkedInProfile, IsActive",
        "Sample Data \u2014 ALUMNI (5 records)",
        [
            "AlumniID", "FirstName", "LastName", "Email", "DateOfBirth", "Gender",
            "Address_City", "Address_State", "Address_PinCode", "GraduationYear",
            "DeptID", "BatchID", "CompanyID", "CurrentPosition", "LinkedInProfile", "IsActive",
        ],
        [
            ["1", "Aarav", "Mehta", "aarav.mehta@email.com", "2001-04-12", "M", "Chennai", "TN", "600001", "2023", "1", "1", "101", "Software Engineer", "in/aaravmehta", "TRUE"],
            ["2", "Priya", "Nair", "priya.nair@email.com", "2000-11-03", "F", "Bengaluru", "KA", "560001", "2023", "1", "2", "102", "Data Analyst", "in/priyanair", "TRUE"],
            ["3", "Rohan", "Gupta", "rohan.gupta@email.com", "2001-07-25", "M", "Mumbai", "MH", "400001", "2024", "2", "3", "103", "Product Manager", "in/rohangupta", "TRUE"],
            ["4", "Sneha", "Reddy", "sneha.reddy@email.com", "2002-01-19", "F", "Hyderabad", "TS", "500001", "2024", "3", "4", "104", "Systems Engineer", "in/snehareddy", "TRUE"],
            ["5", "Vikram", "Singh", "vikram.singh@email.com", "2002-09-30", "M", "Delhi", "DL", "110001", "2025", "4", "5", "105", "Analyst", "in/vikramsingh", "FALSE"],
        ],
        8,
        # Fixed widths (twips) so all 16 columns fit the 7.1in usable page width
        [560, 620, 620, 1240, 700, 480, 620, 480, 620, 620, 480, 480, 520, 800, 700, 480],
    ),
    (
        "StudentID \u2192 FirstName, LastName, Email, DeptID, EnrollmentYear, CurrentSemester, CGPA",
        "Sample Data \u2014 STUDENT (5 records)",
        ["StudentID", "FirstName", "LastName", "Email", "DeptID", "EnrollmentYear", "CurrentSemester", "CGPA"],
        [
            ["1", "Ananya", "Joshi", "ananya.joshi@email.com", "1", "2024", "4", "8.9"],
            ["2", "Karthik", "Iyer", "karthik.iyer@email.com", "2", "2024", "4", "8.4"],
            ["3", "Meera", "Pillai", "meera.pillai@email.com", "3", "2025", "2", "7.8"],
            ["4", "Arjun", "Desai", "arjun.desai@email.com", "4", "2025", "2", "9.1"],
            ["5", "Divya", "Menon", "divya.menon@email.com", "5", "2023", "7", "8.2"],
        ],
        9,
    ),
    (
        "MentorshipID \u2192 AlumniID, StudentID, StartDate, EndDate, Status, MentorshipArea, Goals",
        "Sample Data \u2014 MENTORSHIP (5 records)",
        ["MentorshipID", "AlumniID", "StudentID", "StartDate", "EndDate", "Status", "MentorshipArea", "Goals"],
        [
            ["1", "1", "1", "2025-01-10", "NULL", "Active", "Data Science", "Place in tech"],
            ["2", "2", "2", "2025-02-01", "NULL", "Active", "Career Guidance", "Resume review"],
            ["3", "3", "3", "2024-08-15", "2025-06-30", "Completed", "Web Development", "Full-stack skills"],
            ["4", "4", "4", "2025-03-20", "NULL", "Active", "Research", "Conference paper"],
            ["5", "5", "5", "2024-07-01", "2025-05-31", "Completed", "Entrepreneurship", "Startup guidance"],
        ],
        9,
    ),
    (
        "EventID \u2192 EventName, EventType, EventDate, Venue, OrganizerID",
        "Sample Data \u2014 EVENT (5 records)",
        ["EventID", "EventName", "EventType", "EventDate", "Venue", "OrganizerID"],
        [
            ["1", "Alumni Reunion 2025", "Reunion", "2025-06-15", "Main Auditorium", "1"],
            ["2", "AI Workshop", "Workshop", "2025-04-20", "Tech Hall", "2"],
            ["3", "Career Seminar", "Seminar", "2025-03-10", "Seminar Hall", "3"],
            ["4", "Startup Networking", "Networking", "2025-05-05", "Innovation Lab", "4"],
            ["5", "Freshers Meet 2025", "Networking", "2025-08-01", "Main Auditorium", "5"],
        ],
        10,
    ),
    (
        "DonationID \u2192 DonorID, Amount, DonationDate, PaymentMethod",
        "Sample Data \u2014 DONATION (5 records)",
        ["DonationID", "DonorID", "Amount", "DonationDate", "PaymentMethod"],
        [
            ["1", "1", "50000", "2025-01-10", "Online"],
            ["2", "2", "25000", "2025-02-14", "Online"],
            ["3", "3", "100000", "2025-03-01", "Check"],
            ["4", "4", "15000", "2025-04-22", "Cash"],
            ["5", "5", "75000", "2025-06-30", "Online"],
        ],
        10,
    ),
    (
        "JobID \u2192 JobTitle, JobType, Salary, PostedBy",
        "Sample Data \u2014 JOB (5 records)",
        ["JobID", "JobTitle", "JobType", "Salary", "PostedBy"],
        [
            ["1", "Software Engineer", "Full-Time", "15-20 LPA", "1"],
            ["2", "Data Analyst", "Full-Time", "10-15 LPA", "2"],
            ["3", "Web Developer Intern", "Internship", "5 LPA", "3"],
            ["4", "Project Manager", "Full-Time", "25-30 LPA", "4"],
            ["5", "ML Engineer", "Contract", "20-25 LPA", "5"],
        ],
        10,
    ),
    (
        "(AlumniID, SkillID) \u2192 no non-key attributes",
        "Sample Data \u2014 ALUMNI_SKILL (5 records)",
        ["AlumniID", "SkillID"],
        [
            ["1", "1"],
            ["1", "3"],
            ["2", "2"],
            ["3", "1"],
            ["5", "3"],
        ],
        10,
    ),
    (
        "(AlumniID, EventID) \u2192 no non-key attributes",
        "Sample Data \u2014 ALUMNI_EVENT (5 records)",
        ["AlumniID", "EventID"],
        [
            ["1", "1"],
            ["1", "2"],
            ["2", "1"],
            ["3", "3"],
            ["4", "4"],
        ],
        10,
    ),
    (
        "(AlumniID, PhoneNumber) \u2192 no non-key attributes",
        "Sample Data \u2014 ALUMNI_PHONE (5 records)",
        ["AlumniID", "PhoneNumber"],
        [
            ["1", "+91-9840012345"],
            ["1", "+91-9940012345"],
            ["2", "+91-9880012345"],
            ["3", "+91-9900012345"],
            ["4", "+91-9790012345"],
        ],
        10,
    ),
    (
        "(StudentID, Email) \u2192 no non-key attributes",
        "Sample Data \u2014 STUDENT_EMAIL (5 records)",
        ["StudentID", "Email"],
        [
            ["1", "ananya.joshi@email.com"],
            ["1", "ananya.j@vit.ac.in"],
            ["2", "karthik.iyer@email.com"],
            ["3", "meera.pillai@email.com"],
            ["4", "arjun.desai@email.com"],
        ],
        10,
    ),
]


def main():
    global doc

    # 1. Backup
    shutil.copy2(DOCX, BACKUP)
    print(f"Backup created: {BACKUP}")

    doc = Document(DOCX)

    # ------------------------------------------------------------------
    # A. Remove COMPANY-JOB from the JOB relation
    # ------------------------------------------------------------------
    # A1. JOB schema line (P22): JOB(JobID (PK), JobTitle, JobType, Salary, CompanyID (FK), PostedBy )
    #     -> JOB(JobID (PK), JobTitle, JobType, Salary, PostedBy)
    idx, p = find_paragraph("JOB(JobID (PK)")
    old_text = p.text
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    # Reconstruct with proper formatting: 'JOB(' and ')' were bold, fields not.
    parts = [("JOB(", True)]
    for i, field in enumerate(["JobID (PK)", "JobTitle", "JobType", "Salary", "PostedBy"]):
        if i > 0:
            parts.append((", ", False))
        parts.append((field, False))
    parts.append((")", True))
    for text, bold in parts:
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.bold = bold
    print(f"[A1] P{idx} JOB schema -> {p.text} (was: {old_text})")

    # A2. COMPANY intro
    idx, p = find_paragraph("Stores company information associated with alumni and jobs.")
    p.runs[0].text = "Stores company information associated with alumni."
    print(f"[A2] P{idx} COMPANY intro updated")

    # A3. JOB intro
    idx, p = find_paragraph("Stores job opportunities posted by alumni and associated with companies.")
    p.runs[0].text = "Stores job opportunities posted by alumni."
    print(f"[A3] P{idx} JOB intro updated")

    # A4. JOB FD
    idx, p = find_paragraph("JobID \u2192 JobTitle, JobType, Salary, CompanyID, PostedBy")
    for r in p.runs:
        if "CompanyID" in r.text:
            r.text = "JobID \u2192 JobTitle, JobType, Salary, PostedBy"
    print(f"[A4] P{idx} JOB FD updated -> {p.text}")

    # A5. Delete the bullet "CompanyID represents COMPANY \u2192 JOB."
    idx, p = find_paragraph("CompanyID represents COMPANY")
    delete_paragraph(p)
    print(f"[A5] P{idx} bullet deleted")

    # A6. JOB check table (T12) 3NF row
    for ti, table in enumerate(doc.tables):
        for row in table.rows:
            if len(row.cells) < 3:
                continue
            if "CompanyID and PostedBy are relationship-derived" in row.cells[2].text:
                row.cells[2].text = (
                    "PostedBy is a relationship-derived FK; "
                    "alumni details are stored in their own relation."
                )
                print(f"[A6] T{ti} 3NF reasoning updated")

    # A7. Conclusion: append no COMPANY-JOB relationship
    idx, p = find_paragraph("The corrected report intentionally excludes")
    p.runs[0].text = (
        "The corrected report intentionally excludes structures that are not present in "
        "the finalized ER diagram. In particular, there is no PERSON supertype, no recursive "
        "Alumni-to-Alumni Mentor/Mentee relationship, no JOB_APPLICATION relation, no "
        "EVENT_TYPE lookup relation, no COMPANY_SIZE lookup relation, no MENTORSHIP_AREA "
        "lookup relation, and no COMPANY\u2013JOB relationship."
    )
    print(f"[A7] P{idx} conclusion updated")

    # ------------------------------------------------------------------
    # B. Add sample-data tables after each Functional Dependency paragraph
    # ------------------------------------------------------------------
    for anchor_text, caption, headers, rows, font_size, *rest in SAMPLE_DATA:
        col_widths = rest[0] if rest else None
        idx, p = find_paragraph(anchor_text)
        insert_after(p, caption, headers, rows, font_size, col_widths)
        print(f"[B] Sample table inserted after P{idx} ({caption})")

    # ------------------------------------------------------------------
    # C. Fix relation count 13 -> 14 (add STUDENT_EMAIL to summary)
    # ------------------------------------------------------------------
    # C1. Summary table: add STUDENT_EMAIL row after ALUMNI_PHONE
    for table in doc.tables:
        header_row = [c.text.strip() for c in table.rows[0].cells]
        if header_row == ["Table", "1NF", "2NF", "3NF", "BCNF", "Type"]:
            # find ALUMNI_PHONE row
            for ri, row in enumerate(table.rows):
                if row.cells[0].text.strip() == "ALUMNI_PHONE":
                    new_row = table.add_row()
                    vals = ["STUDENT_EMAIL", "\u2713", "\u2713", "\u2713", "\u2713", "Multivalued attribute handler"]
                    for j, val in enumerate(vals):
                        new_row.cells[j].text = val
                    # Move new row to right after ALUMNI_PHONE
                    alum_phone_tr = table.rows[ri]._tr
                    alum_phone_tr.addnext(new_row._tr)
                    print(f"[C1] STUDENT_EMAIL row added to summary table after ALUMNI_PHONE")
                    break
            break

    # C2. Final count note
    for table in doc.tables:
        if len(table.rows) == 1 and len(table.columns) == 1:
            txt = table.rows[0].cells[0].text
            if "Final count" in txt:
                table.rows[0].cells[0].text = (
                    "Final count  The finalized relational design contains 14 relations: "
                    "10 entity relations plus 4 supporting relations (two M:N junction "
                    "relations and two multivalued-attribute relations)."
                )
                print("[C2] Final count note updated to 14")

    # C3. Final result note
    for table in doc.tables:
        if len(table.rows) == 1 and len(table.columns) == 1:
            txt = table.rows[0].cells[0].text
            if "Final result" in txt:
                table.rows[0].cells[0].text = (
                    "Final result  14 relations are represented in the normalized design, "
                    "and all are in BCNF under the functional dependencies documented in "
                    "this report. No further normalization is required."
                )
                print("[C3] Final result note updated to 14")

    doc.save(DOCX)
    print(f"\nSaved: {DOCX}")


if __name__ == "__main__":
    main()
